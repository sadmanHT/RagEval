"""PDF, DOCX, and HTML loaders that emit normalized RAG-Eval elements only."""

from __future__ import annotations

import html
import re
from collections.abc import Iterable, Sequence
from pathlib import Path
from typing import Any, Protocol

import pymupdf  # type: ignore[import-untyped]
from bs4 import BeautifulSoup, Tag  # type: ignore[import-untyped]
from docx import Document as OpenDocument  # type: ignore[import-untyped]
from docx.table import Table  # type: ignore[import-untyped]
from docx.text.paragraph import Paragraph  # type: ignore[import-untyped]

from rageval.core.errors import DocumentParseError, OCRUnavailableError, UnsupportedSourceError
from rageval.core.ids import fingerprint_mapping, make_element_id
from rageval.corpus.models import CorpusDocument
from rageval.ingestion.models import (
    BatchParseResult,
    OCRMode,
    ParsedDocument,
    ParseFailure,
    ParserConfig,
)
from rageval.models import DocumentElement, DocumentRecord, ElementType, SourceType


class OCRAdapter(Protocol):
    """Stable OCR boundary so deterministic tests need no external service."""

    @property
    def name(self) -> str:
        """Human-readable adapter identifier."""
        ...

    def extract_page_text(self, path: Path, page_number: int, *, dpi: int) -> str:
        """Extract text for a one-based PDF page number."""
        ...


class TesseractOCRAdapter:
    """Local OCR adapter using PyMuPDF rendering plus the Tesseract binary."""

    @property
    def name(self) -> str:
        return "tesseract"

    def extract_page_text(self, path: Path, page_number: int, *, dpi: int) -> str:
        try:
            import pytesseract  # type: ignore[import-untyped]
            from PIL import Image  # type: ignore[import-untyped]

            with pymupdf.open(path) as document:
                page = document.load_page(page_number - 1)
                pixmap = page.get_pixmap(dpi=dpi, alpha=False)
                image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
                return str(pytesseract.image_to_string(image)).strip()
        except (OSError, RuntimeError, ValueError) as exc:
            raise OCRUnavailableError(f"OCR failed for {path} page {page_number}: {exc}") from exc
        except Exception as exc:  # pragma: no cover - dependency wrapper
            module = exc.__class__.__module__
            if module.startswith("pytesseract"):
                raise OCRUnavailableError(
                    f"OCR unavailable for {path} page {page_number}: {exc}"
                ) from exc
            raise


def _config_fingerprint(config: ParserConfig) -> str:
    return fingerprint_mapping(config.model_dump(mode="json"))


def _element(
    *,
    document_id: str,
    ordinal: int,
    kind: ElementType,
    text: str,
    page_number: int | None,
    config_fingerprint: str,
    parser_name: str,
    parser_version: str,
    metadata: dict[str, object] | None = None,
) -> DocumentElement:
    element_metadata: dict[str, object] = {
        "ordinal": ordinal,
        "parser_name": parser_name,
        "parser_version": parser_version,
        "config_fingerprint": config_fingerprint,
    }
    if metadata:
        element_metadata.update(metadata)
    return DocumentElement(
        element_id=make_element_id(
            document_id=document_id,
            ordinal=ordinal,
            kind=kind.value,
            text=text,
            page_number=page_number,
            config_fingerprint=config_fingerprint,
        ),
        document_id=document_id,
        kind=kind,
        text=text,
        page_number=page_number,
        metadata=element_metadata,
    )


def _classify_text(
    text: str,
    *,
    heading: bool = False,
    top: bool = False,
    bottom: bool = False,
) -> ElementType:
    stripped = text.strip()
    if heading:
        return ElementType.TITLE
    if top:
        return ElementType.HEADER
    if bottom:
        return ElementType.FOOTER
    if re.match(r"^(figure|table)\s+\d+[\s:.-]", stripped, flags=re.IGNORECASE):
        return ElementType.CAPTION
    if re.match(r"^(?:[-*•]\s+|\d+[.)]\s+|[A-Za-z][.)]\s+)", stripped):
        return ElementType.LIST_ITEM
    return ElementType.TEXT


def _bbox_intersects(
    candidate: tuple[float, float, float, float],
    table_boxes: Sequence[tuple[float, float, float, float]],
) -> bool:
    x0, y0, x1, y1 = candidate
    for tx0, ty0, tx1, ty1 in table_boxes:
        overlap_x = max(0.0, min(x1, tx1) - max(x0, tx0))
        overlap_y = max(0.0, min(y1, ty1) - max(y0, ty0))
        if overlap_x * overlap_y > 0:
            return True
    return False


def _table_text(rows: Sequence[Sequence[object | None]]) -> str:
    normalized: list[str] = []
    for row in rows:
        normalized.append("\t".join("" if cell is None else str(cell).strip() for cell in row))
    return "\n".join(normalized).strip()


def _table_html(rows: Sequence[Sequence[object | None]]) -> str:
    body: list[str] = ["<table>"]
    for row in rows:
        body.append("<tr>")
        for cell in row:
            value = "" if cell is None else str(cell).strip()
            body.append(f"<td>{html.escape(value)}</td>")
        body.append("</tr>")
    body.append("</table>")
    return "".join(body)


def _pdf_text_items(
    page: Any,
    *,
    page_index: int,
    table_boxes: Sequence[tuple[float, float, float, float]],
    section_hint: str | None,
) -> tuple[list[tuple[float, int, ElementType, str, dict[str, object]]], str | None]:
    page_items: list[tuple[float, int, ElementType, str, dict[str, object]]] = []
    page_height = float(page.rect.height)
    blocks = page.get_text("dict", sort=True).get("blocks", [])
    current_section = section_hint
    for block_index, block in enumerate(blocks):
        if int(block.get("type", 0)) != 0:
            continue
        bbox_raw = block.get("bbox")
        if not isinstance(bbox_raw, (list, tuple)) or len(bbox_raw) != 4:
            raise DocumentParseError(f"unexpected PDF text block shape on page {page_index + 1}")
        bbox = tuple(float(value) for value in bbox_raw)
        if _bbox_intersects(bbox, table_boxes):
            continue
        line_texts: list[str] = []
        font_sizes: list[float] = []
        bold = False
        for line in block.get("lines", []):
            for span in line.get("spans", []):
                span_text = str(span.get("text", ""))
                if span_text:
                    line_texts.append(span_text)
                size = span.get("size")
                if isinstance(size, (int, float)):
                    font_sizes.append(float(size))
                flags = span.get("flags")
                if isinstance(flags, int) and flags & 16:
                    bold = True
        text = " ".join(part.strip() for part in line_texts if part.strip()).strip()
        if not text:
            continue
        heading = bool(font_sizes and max(font_sizes) >= 14.0 and len(text) <= 180) or (
            bold and len(text) <= 100
        )
        kind = _classify_text(
            text,
            heading=heading,
            top=bbox[1] <= page_height * 0.06 and not heading,
            bottom=bbox[3] >= page_height * 0.94 and not heading,
        )
        if kind is ElementType.TITLE:
            current_section = text
        page_items.append(
            (
                bbox[1],
                1,
                kind,
                text,
                {
                    "source": "native",
                    "source_offset": {
                        "page_index": page_index,
                        "block_index": block_index,
                    },
                    "bbox": list(bbox),
                    "section_hint": current_section,
                },
            )
        )
    return page_items, current_section


def _pdf_table_items(
    page: Any,
    *,
    page_index: int,
    section_hint: str | None,
    detect_tables: bool,
) -> tuple[
    list[tuple[float, int, ElementType, str, dict[str, object]]],
    list[tuple[float, float, float, float]],
]:
    page_items: list[tuple[float, int, ElementType, str, dict[str, object]]] = []
    table_boxes: list[tuple[float, float, float, float]] = []
    if not detect_tables or not hasattr(page, "find_tables"):
        return page_items, table_boxes
    try:
        found = page.find_tables()
        for table_index, table in enumerate(found.tables):
            bbox = tuple(float(value) for value in table.bbox)
            table_boxes.append(bbox)
            rows = table.extract()
            text = _table_text(rows)
            if not text:
                continue
            page_items.append(
                (
                    bbox[1],
                    0,
                    ElementType.TABLE,
                    text,
                    {
                        "source": "native",
                        "source_offset": {
                            "page_index": page_index,
                            "table_index": table_index,
                        },
                        "bbox": list(bbox),
                        "table_html": _table_html(rows),
                        "section_hint": section_hint,
                    },
                )
            )
    except (AttributeError, RuntimeError, ValueError):
        # Table detection is additive. Text extraction still proceeds and no page is dropped.
        return [], []
    return page_items, table_boxes


def _parse_pdf(
    path: Path,
    document: DocumentRecord,
    *,
    config: ParserConfig,
    ocr_adapter: OCRAdapter | None,
) -> ParsedDocument:
    parser_name = "pymupdf"
    fingerprint = _config_fingerprint(config)
    raw: list[tuple[ElementType, str, int | None, dict[str, object]]] = []
    used_ocr = False
    section_hint: str | None = None

    try:
        pdf = pymupdf.open(path)
    except Exception as exc:
        raise DocumentParseError(f"cannot open PDF {path}: {exc}") from exc

    page_count = pdf.page_count
    try:
        if pdf.needs_pass:
            raise DocumentParseError(f"encrypted PDF requires a password: {path}")
        for page_index in range(page_count):
            page_number = page_index + 1
            page = pdf.load_page(page_index)
            native_text = str(page.get_text("text", sort=True)).strip()
            table_items, table_boxes = _pdf_table_items(
                page,
                page_index=page_index,
                section_hint=section_hint,
                detect_tables=config.detect_tables,
            )
            text_items, section_hint = _pdf_text_items(
                page,
                page_index=page_index,
                table_boxes=table_boxes,
                section_hint=section_hint,
            )
            page_items = table_items + text_items

            should_ocr = config.ocr_mode is OCRMode.ALWAYS or (
                config.ocr_mode is OCRMode.FALLBACK
                and len(native_text) < config.ocr_min_native_chars
            )
            if should_ocr:
                if ocr_adapter is None:
                    raise OCRUnavailableError(
                        "OCR is required for "
                        f"{path} page {page_number} but no adapter is configured"
                    )
                ocr_text = ocr_adapter.extract_page_text(
                    path,
                    page_number,
                    dpi=config.ocr_dpi,
                ).strip()
                if ocr_text:
                    used_ocr = True
                    page_items.append(
                        (
                            0.0,
                            2,
                            ElementType.OCR_TEXT,
                            ocr_text,
                            {
                                "source": "ocr",
                                "ocr_adapter": ocr_adapter.name,
                                "source_offset": {"page_index": page_index},
                                "section_hint": section_hint,
                            },
                        )
                    )
                elif not native_text:
                    raise DocumentParseError(
                        f"page {page_number} of {path} produced no native text and empty OCR"
                    )

            page_items.sort(key=lambda item: (item[0], item[1], item[3]))
            raw.extend(
                (kind, text, page_number, metadata) for _, _, kind, text, metadata in page_items
            )
            if config.emit_page_breaks and page_index < page_count - 1:
                raw.append(
                    (
                        ElementType.PAGE_BREAK,
                        "",
                        page_number,
                        {
                            "source": "synthetic",
                            "source_offset": {"after_page": page_number},
                            "section_hint": section_hint,
                        },
                    )
                )
    except (DocumentParseError, OCRUnavailableError):
        raise
    except Exception as exc:
        raise DocumentParseError(f"failed while parsing PDF {path}: {exc}") from exc
    finally:
        pdf.close()

    elements = tuple(
        _element(
            document_id=document.document_id,
            ordinal=ordinal,
            kind=kind,
            text=text,
            page_number=page,
            config_fingerprint=fingerprint,
            parser_name=parser_name,
            parser_version=config.parser_version,
            metadata=metadata,
        )
        for ordinal, (kind, text, page, metadata) in enumerate(raw)
    )
    return ParsedDocument(
        document=document,
        parser_name=parser_name,
        parser_version=config.parser_version,
        config_fingerprint=fingerprint,
        elements=elements,
        used_ocr=used_ocr,
        metadata={"path": path.name, "page_count": page_count},
    )


def _iter_docx_blocks(document: Any) -> Iterable[Paragraph | Table]:
    paragraphs = iter(document.paragraphs)
    tables = iter(document.tables)
    paragraph = next(paragraphs, None)
    table = next(tables, None)
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            if paragraph is not None:
                yield paragraph
                paragraph = next(paragraphs, None)
        elif child.tag.endswith("}tbl"):
            if table is not None:
                yield table
                table = next(tables, None)


def _docx_has_page_break(paragraph: Paragraph) -> bool:
    xml = str(paragraph._p.xml)  # type: ignore[attr-defined]
    return 'w:type="page"' in xml or "lastRenderedPageBreak" in xml


def _parse_docx(path: Path, document: DocumentRecord, *, config: ParserConfig) -> ParsedDocument:
    parser_name = "python-docx"
    fingerprint = _config_fingerprint(config)
    raw: list[tuple[ElementType, str, int | None, dict[str, object]]] = []
    page_number = 1
    section_hint: str | None = None
    paragraph_index = 0
    table_index = 0

    try:
        doc = OpenDocument(path)
        for block in _iter_docx_blocks(doc):
            if isinstance(block, Paragraph):
                text = block.text.strip()
                style_name = block.style.name if block.style is not None else ""
                if text:
                    heading = style_name.lower().startswith(("heading", "title"))
                    if heading:
                        kind = ElementType.TITLE
                        section_hint = text
                    elif style_name.lower().startswith("list"):
                        kind = ElementType.LIST_ITEM
                    elif style_name.lower().startswith("caption"):
                        kind = ElementType.CAPTION
                    else:
                        kind = _classify_text(text)
                    raw.append(
                        (
                            kind,
                            text,
                            page_number,
                            {
                                "source": "native",
                                "style": style_name,
                                "source_offset": {"paragraph_index": paragraph_index},
                                "section_hint": section_hint,
                            },
                        )
                    )
                if config.emit_page_breaks and _docx_has_page_break(block):
                    raw.append(
                        (
                            ElementType.PAGE_BREAK,
                            "",
                            page_number,
                            {
                                "source": "native",
                                "source_offset": {"paragraph_index": paragraph_index},
                                "section_hint": section_hint,
                            },
                        )
                    )
                    page_number += 1
                paragraph_index += 1
            elif isinstance(block, Table):
                rows = [[cell.text for cell in row.cells] for row in block.rows]
                text = _table_text(rows)
                if text:
                    raw.append(
                        (
                            ElementType.TABLE,
                            text,
                            page_number,
                            {
                                "source": "native",
                                "source_offset": {"table_index": table_index},
                                "table_html": _table_html(rows),
                                "section_hint": section_hint,
                            },
                        )
                    )
                table_index += 1

        for section_index, section in enumerate(doc.sections):
            parts = (
                (ElementType.HEADER, section.header.paragraphs),
                (ElementType.FOOTER, section.footer.paragraphs),
            )
            for kind, paragraphs in parts:
                for part_index, paragraph in enumerate(paragraphs):
                    text = paragraph.text.strip()
                    if text:
                        raw.append(
                            (
                                kind,
                                text,
                                None,
                                {
                                    "source": "native",
                                    "source_offset": {
                                        "section_index": section_index,
                                        "part_index": part_index,
                                    },
                                    "section_hint": section_hint,
                                },
                            )
                        )
    except Exception as exc:
        raise DocumentParseError(f"cannot parse DOCX {path}: {exc}") from exc

    elements = tuple(
        _element(
            document_id=document.document_id,
            ordinal=ordinal,
            kind=kind,
            text=text,
            page_number=page,
            config_fingerprint=fingerprint,
            parser_name=parser_name,
            parser_version=config.parser_version,
            metadata=metadata,
        )
        for ordinal, (kind, text, page, metadata) in enumerate(raw)
    )
    return ParsedDocument(
        document=document,
        parser_name=parser_name,
        parser_version=config.parser_version,
        config_fingerprint=fingerprint,
        elements=elements,
        metadata={"path": path.name},
    )


def _is_page_break(tag: Tag) -> bool:
    style = str(tag.get("style", "")).replace(" ", "").lower()
    classes = " ".join(str(value).lower() for value in tag.get("class", []))
    return (tag.name == "hr" and "page-break" in classes) or (
        "page-break-before:always" in style or "page-break-after:always" in style
    )


def _parse_html(path: Path, document: DocumentRecord, *, config: ParserConfig) -> ParsedDocument:
    parser_name = "beautifulsoup"
    fingerprint = _config_fingerprint(config)
    raw: list[tuple[ElementType, str, int | None, dict[str, object]]] = []
    section_hint: str | None = None

    try:
        soup = BeautifulSoup(path.read_text(encoding="utf-8"), "html.parser")
    except (OSError, UnicodeError, ValueError) as exc:
        raise DocumentParseError(f"cannot parse HTML {path}: {exc}") from exc

    for unwanted in soup(["script", "style", "noscript"]):
        unwanted.decompose()

    tags = soup.find_all(
        ["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "table", "header", "footer", "hr"]
    )
    for source_index, node in enumerate(tags):
        if not isinstance(node, Tag):
            continue
        parent = node.find_parent(["table", "li", "header", "footer"])
        if parent is not None and node.name not in {"table", "li", "header", "footer"}:
            continue
        if _is_page_break(node):
            if config.emit_page_breaks:
                raw.append(
                    (
                        ElementType.PAGE_BREAK,
                        "",
                        None,
                        {
                            "source": "native",
                            "source_offset": {"dom_index": source_index},
                            "section_hint": section_hint,
                        },
                    )
                )
            continue

        text = node.get_text(" ", strip=True)
        if not text:
            continue
        if node.name and node.name.startswith("h"):
            kind = ElementType.TITLE
            section_hint = text
        elif node.name == "li":
            kind = ElementType.LIST_ITEM
        elif node.name == "table":
            kind = ElementType.TABLE
        elif node.name == "header":
            kind = ElementType.HEADER
        elif node.name == "footer":
            kind = ElementType.FOOTER
        else:
            classes = {str(value).lower() for value in node.get("class", [])}
            kind = (
                ElementType.CAPTION
                if "caption" in classes or re.match(r"^(figure|table)\s+\d+", text, re.I)
                else ElementType.TEXT
            )

        metadata: dict[str, object] = {
            "source": "native",
            "source_offset": {"dom_index": source_index},
            "tag": node.name or "",
            "section_hint": section_hint,
        }
        if kind is ElementType.TABLE:
            rows = [
                [
                    cell.get_text(" ", strip=True)
                    for cell in row.find_all(["th", "td"], recursive=False)
                ]
                for row in node.find_all("tr")
            ]
            metadata["table_html"] = str(node)
            normalized = _table_text(rows)
            if normalized:
                text = normalized
        raw.append((kind, text, None, metadata))

    elements = tuple(
        _element(
            document_id=document.document_id,
            ordinal=ordinal,
            kind=kind,
            text=text,
            page_number=page,
            config_fingerprint=fingerprint,
            parser_name=parser_name,
            parser_version=config.parser_version,
            metadata=metadata,
        )
        for ordinal, (kind, text, page, metadata) in enumerate(raw)
    )
    return ParsedDocument(
        document=document,
        parser_name=parser_name,
        parser_version=config.parser_version,
        config_fingerprint=fingerprint,
        elements=elements,
        metadata={"path": path.name},
    )


def parse_document(
    path: Path,
    document: DocumentRecord,
    *,
    config: ParserConfig | None = None,
    ocr_adapter: OCRAdapter | None = None,
) -> ParsedDocument:
    """Parse one supported file into project-owned normalized contracts."""

    resolved = path.resolve()
    if not resolved.is_file():
        raise DocumentParseError(f"source file does not exist: {path}")
    active = config or ParserConfig()
    suffix = resolved.suffix.lower()
    expected_suffixes = {
        SourceType.PDF: {".pdf"},
        SourceType.DOCX: {".docx"},
        SourceType.HTML: {".html", ".htm"},
    }
    if suffix not in expected_suffixes[document.source_type]:
        raise UnsupportedSourceError(
            f"source type {document.source_type.value} does not match file extension {suffix!r}"
        )

    if document.source_type is SourceType.PDF:
        adapter = ocr_adapter
        if adapter is None and active.ocr_mode is not OCRMode.DISABLED:
            adapter = TesseractOCRAdapter()
        return _parse_pdf(resolved, document, config=active, ocr_adapter=adapter)
    if document.source_type is SourceType.DOCX:
        return _parse_docx(resolved, document, config=active)
    if document.source_type is SourceType.HTML:
        return _parse_html(resolved, document, config=active)
    raise UnsupportedSourceError(f"unsupported source type: {document.source_type}")


def parse_corpus_document(
    item: CorpusDocument,
    root: Path,
    *,
    config: ParserConfig | None = None,
    ocr_adapter: OCRAdapter | None = None,
) -> ParsedDocument:
    """Resolve one manifest item under its corpus root and parse it."""

    root = root.resolve()
    path = (root / item.relative_path).resolve()
    try:
        path.relative_to(root)
    except ValueError as exc:
        raise DocumentParseError(f"source path escapes corpus root: {item.relative_path}") from exc
    return parse_document(path, item.record, config=config, ocr_adapter=ocr_adapter)


def parse_batch(
    items: Sequence[CorpusDocument],
    root: Path,
    *,
    config: ParserConfig | None = None,
    ocr_adapter: OCRAdapter | None = None,
) -> BatchParseResult:
    """Parse a batch while retaining per-file failures instead of aborting on first error."""

    successes: list[ParsedDocument] = []
    failures: list[ParseFailure] = []
    for item in items:
        try:
            successes.append(
                parse_corpus_document(item, root, config=config, ocr_adapter=ocr_adapter)
            )
        except (DocumentParseError, OCRUnavailableError, UnsupportedSourceError) as exc:
            failures.append(
                ParseFailure(
                    document_id=item.record.document_id,
                    relative_path=item.relative_path,
                    error_type=type(exc).__name__,
                    message=str(exc),
                )
            )
    return BatchParseResult(successes=tuple(successes), failures=tuple(failures))
